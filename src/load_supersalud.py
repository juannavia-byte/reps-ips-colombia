"""
Carga los estados financieros de Supersalud y los cruza contra el REPS por NIT.

Uso:  python src/load_supersalud.py --dsn ... --raw data/raw

Reglas del cruce
----------------
· Se cruza por (tipo_identificacion='NI', numero_identificacion=NIT). Las IPS
  públicas traen además código de habilitación, que se guarda pero no se usa
  como llave: el NIT es la única llave común a los dos archivos.
· Una fila financiera SIN contraparte en REPS se carga igual, con
  prestador_id NULL. Así la tasa de cruce es medible en vez de asumida.
· Ningún prestador del REPS se elimina por no tener financiero.

Estructura de los archivos (verificada)
---------------------------------------
  privadas  FT001-01 cabecera fila 14, 390 cols · FT001-02 fila 14, 351 · FT001-03 fila 15, 266
  públicas  IPS_PUBLICAS_DIC_2021 cabecera fila 13, 167 cols

Los tres grupos NIIF usan planes de cuentas distintos, así que el detalle va a
JSONB y solo se promueven a columnas las cinco cifras comparables.
"""
from __future__ import annotations

import argparse
import json
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import psycopg

# nombre lógico -> (nombres exactos aceptados, prefijos de respaldo)
AGREGADOS = {
    "activos":    (("ACTIVOS",),    ("1 ACTIVOS",)),
    "pasivos":    (("PASIVOS",),    ("2 PASIVOS",)),
    "patrimonio": (("PATRIMONIO",), ("3 PATRIMONIO",)),
    "ingresos":   (("INGRESOS",),   ("4 INGRESOS",)),
    "costos":     (("COSTOS",),     ("6 COSTOS",)),
}

HOJAS_PRIVADAS = {"FT001-01": "privada_g1", "FT001-02": "privada_g2", "FT001-03": "privada_g3"}


def cabecera(ws, limite=30):
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=limite, values_only=True), start=1):
        vals = [str(c).strip() if c is not None else "" for c in row]
        if any(v.upper() in ("NIT", "CÓDIGO HABILITACIÓN") for v in vals[:4]):
            return i, vals
    return None, None


def resolver(cols: list[str]) -> dict[str, int]:
    """Ubica las columnas agregadas, tolerando que cambien de nombre entre grupos."""
    arriba = [c.strip().upper() for c in cols]
    salida = {}
    for logico, (exactos, prefijos) in AGREGADOS.items():
        idx = next((i for i, c in enumerate(arriba) if c in exactos), None)
        if idx is None:
            idx = next((i for i, c in enumerate(arriba)
                        if any(c.startswith(p) for p in prefijos)), None)
        if idx is not None:
            salida[logico] = idx
    return salida


def numero(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    try:
        return Decimal(re.sub(r"[^\d.\-]", "", str(v)) or "0")
    except InvalidOperation:
        return None


def solo_digitos(v) -> str | None:
    if v is None:
        return None
    s = re.sub(r"\D", "", str(v))
    return s or None


def filas_hoja(ws, fila_cab: int, cols: list[str], col_nit: int, col_razon: int | None):
    for row in ws.iter_rows(min_row=fila_cab + 1, values_only=True):
        nit = solo_digitos(row[col_nit] if col_nit < len(row) else None)
        if not nit:
            continue
        razon = None
        if col_razon is not None and col_razon < len(row) and row[col_razon] is not None:
            razon = str(row[col_razon]).strip() or None
        yield nit, razon, row


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dsn", required=True)
    ap.add_argument("--raw", default="data/raw")
    args = ap.parse_args()
    raw = Path(args.raw)
    man = json.loads(sorted(raw.glob("manifiesto_supersalud_*.json"))[-1].read_text())
    por_clave = {a["clave"]: a for a in man["archivos"]}

    cx = psycopg.connect(args.dsn, autocommit=False)
    resumen = {}

    # Índice NIT -> prestador_id (solo personas jurídicas).
    with cx.cursor() as cur:
        cur.execute("SELECT numero_identificacion, id FROM reps.prestador "
                    "WHERE tipo_identificacion = 'NI'")
        por_nit = dict(cur.fetchall())
    print(f"· índice REPS: {len(por_nit):,} NIT de personas jurídicas", flush=True)

    for clave, meta in por_clave.items():
        ruta = Path(meta["archivo"])
        with cx.cursor() as cur:
            cur.execute(
                """INSERT INTO reps.extraccion
                   (fuente, recurso, url, extraido_en, fecha_corte_declarada, archivo,
                    sha256, tamano_bytes, filas_crudas, separador, notas)
                   VALUES ('supersalud',%s,%s,%s,%s,%s,%s,%s,NULL,NULL,%s) RETURNING id""",
                (clave, meta["url"], man["extraido_en"], f"31-dic-{meta['vigencia']}",
                 str(ruta), meta["sha256"], meta["tamano_bytes"],
                 "última vigencia publicada por Supersalud; no existen cortes posteriores"),
            )
            ext_id = cur.fetchone()[0]
        cx.commit()

        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        for ws in wb.worksheets:
            origen = HOJAS_PRIVADAS.get(ws.title, "ese" if meta["naturaleza"] == "publica" else None)
            if origen is None:
                continue
            fila_cab, cols = cabecera(ws)
            if not fila_cab:
                continue
            arriba = [c.strip().upper() for c in cols]
            col_nit = arriba.index("NIT")
            col_razon = next((i for i, c in enumerate(arriba)
                              if c in ("RAZÓN SOCIAL", "RAZON SOCIAL", "ENTIDAD")), None)
            col_hab = next((i for i, c in enumerate(arriba) if "HABILITACI" in c), None)
            col_niv = next((i for i, c in enumerate(arriba) if c == "NIVEL"), None)
            col_dep = next((i for i, c in enumerate(arriba) if c == "DEPARTAMENTO"), None)
            col_mun = next((i for i, c in enumerate(arriba) if c == "MUNICIPIO"), None)
            agg = resolver(cols)

            n = cruzados = 0
            with cx.cursor() as cur:
                for nit, razon, row in filas_hoja(ws, fila_cab, cols, col_nit, col_razon):
                    pid = por_nit.get(nit)
                    cuentas = {
                        cols[i]: float(x)
                        for i, x in enumerate(row)
                        if i < len(cols) and cols[i] and isinstance(x, (int, float))
                    }
                    vals = {k: numero(row[i]) if i < len(row) else None for k, i in agg.items()}
                    cur.execute(
                        """INSERT INTO reps.prestador_financiero
                           (prestador_id, nit_reportado, razon_social_reportada, vigencia, origen,
                            codigo_habilitacion_rep, nivel, departamento, municipio, cuentas,
                            activos, pasivos, patrimonio, ingresos, costos, extraccion_id)
                           VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                           ON CONFLICT (nit_reportado, vigencia, origen) DO NOTHING""",
                        (pid, nit, razon, meta["vigencia"], origen,
                         str(row[col_hab]).strip() if col_hab is not None and col_hab < len(row) and row[col_hab] else None,
                         str(row[col_niv]).strip() if col_niv is not None and col_niv < len(row) and row[col_niv] else None,
                         str(row[col_dep]).strip() if col_dep is not None and col_dep < len(row) and row[col_dep] else None,
                         str(row[col_mun]).strip() if col_mun is not None and col_mun < len(row) and row[col_mun] else None,
                         json.dumps(cuentas, ensure_ascii=False),
                         vals.get("activos"), vals.get("pasivos"), vals.get("patrimonio"),
                         vals.get("ingresos"), vals.get("costos"), ext_id),
                    )
                    n += 1
                    cruzados += pid is not None
                    if pid is None:
                        cur.execute(
                            """INSERT INTO reps.qa_inconsistencia
                               (extraccion_id, tabla, clave, tipo, detalle, valor_crudo)
                               VALUES (%s,'prestador_financiero',%s,'sin_match_reps',%s,%s)""",
                            (ext_id, nit,
                             "NIT reportado a Supersalud que no aparece como persona jurídica "
                             "en el REPS vigente (IPS cerrada, fusionada o de otro régimen)",
                             (razon or "")[:500]),
                        )
            cx.commit()
            resumen[f"{clave}:{ws.title}"] = {
                "origen": origen, "filas": n, "cruzados": cruzados,
                "sin_match": n - cruzados,
                "tasa_cruce": round(100 * cruzados / n, 2) if n else 0,
                "agregados_detectados": sorted(agg),
                "cabecera_fila": fila_cab,
            }
            print(f"  {clave}/{ws.title}: {n:,} filas · {cruzados:,} cruzadas "
                  f"({resumen[f'{clave}:{ws.title}']['tasa_cruce']}%)", flush=True)
        wb.close()

    cx.close()
    (raw.parent / "resumen_supersalud.json").write_text(
        json.dumps(resumen, indent=2, ensure_ascii=False))
    print("\n Supersalud cargado.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
